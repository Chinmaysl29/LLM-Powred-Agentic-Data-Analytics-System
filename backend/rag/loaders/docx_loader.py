"""DOCX Document Loader for RAG ingestion.

Uses python-docx to extract paragraphs, headings, bullet lists,
tables, and core document metadata.
"""

import io
import logging
from pathlib import Path
from typing import Any
from uuid import uuid4

import docx

from backend.app.schemas.document_loader import LoadedDocument
from backend.rag.loaders.base_loader import BaseDocumentLoader

logger = logging.getLogger(__name__)


class DocxLoader(BaseDocumentLoader):
    """Loader extracting structured text from Microsoft Word (.docx) documents."""

    def supported_extensions(self) -> list[str]:
        return [".docx"]

    def load(
        self,
        file_input: str | Path | bytes,
        filename: str | None = None,
        **kwargs: Any,
    ) -> LoadedDocument:
        raw_bytes, resolved_name = self._read_bytes_and_filename(file_input, filename)
        logger.info("Loading DOCX document '%s' (%d bytes)", resolved_name, len(raw_bytes))

        stream = io.BytesIO(raw_bytes)
        doc = docx.Document(stream)

        sections: list[str] = []

        # 1. Extract Paragraphs & Headings
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                sections.append(text)

        # 2. Extract Tables as Structured Markdown Tables
        for table in doc.tables:
            table_lines: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                # Avoid completely blank table rows
                if any(cells):
                    table_lines.append("| " + " | ".join(cells) + " |")

            if table_lines:
                # If table has a header row, inject a markdown separator row
                if len(table_lines) > 1:
                    col_count = len(table.rows[0].cells)
                    sep = "| " + " | ".join(["---"] * col_count) + " |"
                    table_lines.insert(1, sep)
                sections.append("\n".join(table_lines))

        # 3. Extract Metadata Properties
        title = None
        author = None
        try:
            core_props = doc.core_properties
            title = core_props.title if core_props and core_props.title else None
            author = core_props.author if core_props and core_props.author else None
        except Exception as prop_exc:
            logger.debug("Could not read docx core properties: %s", prop_exc)

        full_text = "\n\n".join(sections)
        normalized_content = self.normalize_text(full_text)

        metadata = self.build_metadata(
            filename=resolved_name,
            file_type="docx",
            size=len(raw_bytes),
            content=normalized_content,
            title=title,
            author=author,
        )

        return LoadedDocument(
            document_id=str(uuid4()),
            content=normalized_content,
            metadata=metadata,
        )
