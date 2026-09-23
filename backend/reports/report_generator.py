"""Enterprise Report Generation Engine for Phase 8.

Supports Formats:
- PDF (via ReportLab)
- DOCX (via python-docx)
- Excel (via openpyxl)
- PPTX (via python-pptx)

Supports Report Types:
- Executive Summary
- Analytics Report
- Forecast Report
- Recommendations Report

Outputs standard format:
{
  "report_id": "rep_102",
  "file_url": "/reports/rep_102.pdf"
}
"""

from __future__ import annotations

import logging
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

DEFAULT_OUTPUT_DIR = Path("data/reports")


class ReportGenerator:
    """Multi-format enterprise report generator."""

    def __init__(self, output_dir: Path | str | None = None) -> None:
        self.output_dir = Path(output_dir or DEFAULT_OUTPUT_DIR)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(
        self,
        report_type: str,
        format_type: str = "pdf",
        title: str | None = None,
        data: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Generate report in requested format and return file descriptor."""
        clean_type = report_type.lower().replace(" ", "_")
        clean_format = format_type.lower().replace(".", "").strip()
        if clean_format == "excel":
            clean_format = "xlsx"

        report_id = f"rep_{uuid.uuid4().hex[:8]}"
        filename = f"{report_id}.{clean_format}"
        file_path = self.output_dir / filename

        report_title = title or clean_type.replace("_", " ").title()
        report_data = data or self._get_default_report_data(clean_type)

        if clean_format == "pdf":
            self._generate_pdf(file_path, report_title, clean_type, report_data)
        elif clean_format == "docx":
            self._generate_docx(file_path, report_title, clean_type, report_data)
        elif clean_format == "xlsx":
            self._generate_excel(file_path, report_title, clean_type, report_data)
        elif clean_format == "pptx":
            self._generate_pptx(file_path, report_title, clean_type, report_data)
        else:
            raise ValueError(f"Unsupported report format: {format_type}")

        file_url = f"/reports/{filename}"

        return {
            "report_id": report_id,
            "file_url": file_url,
            "file_path": str(file_path),
            "format": clean_format,
            "report_type": clean_type,
            "generated_at": datetime.now(timezone.utc).isoformat(),
        }

    def _get_default_report_data(self, report_type: str) -> dict[str, Any]:
        """Provide default structured content for each report type."""
        if "executive" in report_type:
            return {
                "summary": "Enterprise performance has grown 18.5% YoY with healthy unit margins.",
                "kpis": {"Annual Recurring Revenue": "$14.2M", "Gross Margin": "76.4%", "Customer Retention": "94.2%"},
                "insights": [
                    "Operating costs stabilized with 12% server efficiency improvement.",
                    "North American market is outperforming targets by 22%.",
                ],
                "action_items": [
                    "Accelerate high-tier enterprise marketing campaigns.",
                    "Review inventory allocation for top Q3 demand nodes.",
                ],
            }
        elif "forecast" in report_type:
            return {
                "summary": "30-day demand forecast indicates an upward trajectory with 95% confidence interval.",
                "kpis": {"Predicted Mean": 128400, "Lower Bound": 115000, "Upper Bound": 142000, "MAPE": "4.2%"},
                "insights": ["Trend exhibits weekly seasonality with peaks on Thursdays.", "No significant change points detected."],
                "action_items": ["Procure raw inventory 10 days in advance of mid-month spike."],
            }
        elif "recommendation" in report_type:
            return {
                "summary": "Decision intelligence engine generated 4 high-impact action recommendations.",
                "kpis": {"Identified Savings": "$240,000", "Revenue Upside": "$580,000", "ROI": "3.8x"},
                "insights": [
                    "Cloud instance downsizing offers immediate $18k/month recurring savings.",
                    "Dynamic pricing tier adjustment projected to increase conversion by 4.5%.",
                ],
                "action_items": [
                    "Approve cloud instance resizing plan.",
                    "A/B test premium subscription pricing increase by 6%.",
                ],
            }
        else:  # analytics report
            return {
                "summary": "Deep-dive dataset analysis across 45,000 records and 18 attributes.",
                "kpis": {"Total Rows": 45000, "Completeness": "99.8%", "Outliers": 124},
                "insights": ["Strong correlation (r=0.81) between customer engagement score and lifetime value."],
                "action_items": ["Target customer onboarding sessions to raise 7-day engagement."],
            }

    def _generate_pdf(
        self, path: Path, title: str, report_type: str, data: dict[str, Any]
    ) -> None:
        """Generate PDF report using ReportLab."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        elements = []

        # Title
        title_style = ParagraphStyle("ReportTitle", parent=styles["Heading1"], fontSize=20, leading=24, textColor=colors.HexColor("#1e293b"))
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 10))

        # Metadata
        meta_style = ParagraphStyle("Meta", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"))
        elements.append(Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Type: {report_type.upper()}", meta_style))
        elements.append(Spacer(1, 15))

        # Executive Summary
        h2_style = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=14, leading=18, textColor=colors.HexColor("#0f172a"))
        body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14, textColor=colors.HexColor("#334155"))
        elements.append(Paragraph("Executive Overview", h2_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(data.get("summary", ""), body_style))
        elements.append(Spacer(1, 15))

        # KPIs Table
        kpis = data.get("kpis", {})
        if kpis:
            elements.append(Paragraph("Key Metrics", h2_style))
            elements.append(Spacer(1, 6))
            table_data = [["Metric", "Value"]]
            for k, v in kpis.items():
                table_data.append([str(k), str(v)])

            t = Table(table_data, colWidths=[280, 200])
            t.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ])
            )
            elements.append(t)
            elements.append(Spacer(1, 15))

        # Insights & Actions
        insights = data.get("insights", [])
        if insights:
            elements.append(Paragraph("Strategic Insights", h2_style))
            elements.append(Spacer(1, 6))
            for item in insights:
                elements.append(Paragraph(f"• {item}", body_style))
            elements.append(Spacer(1, 15))

        actions = data.get("action_items", [])
        if actions:
            elements.append(Paragraph("Recommended Next Steps", h2_style))
            elements.append(Spacer(1, 6))
            for action in actions:
                elements.append(Paragraph(f"✓ {action}", body_style))

        doc.build(elements)

    def _generate_docx(
        self, path: Path, title: str, report_type: str, data: dict[str, Any]
    ) -> None:
        """Generate DOCX document."""
        import docx

        doc = docx.Document()
        doc.add_heading(title, level=0)
        doc.add_paragraph(f"Report Type: {report_type.upper()} | Generated: {datetime.now(timezone.utc).isoformat()}")

        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(data.get("summary", ""))

        kpis = data.get("kpis", {})
        if kpis:
            doc.add_heading("Key Performance Indicators", level=1)
            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Metric"
            hdr_cells[1].text = "Value"
            for k, v in kpis.items():
                row_cells = table.add_row().cells
                row_cells[0].text = str(k)
                row_cells[1].text = str(v)

        insights = data.get("insights", [])
        if insights:
            doc.add_heading("Key Insights", level=1)
            for item in insights:
                doc.add_paragraph(item, style="List Bullet")

        actions = data.get("action_items", [])
        if actions:
            doc.add_heading("Action Plan", level=1)
            for action in actions:
                doc.add_paragraph(action, style="List Bullet")

        doc.save(str(path))

    def _generate_excel(
        self, path: Path, title: str, report_type: str, data: dict[str, Any]
    ) -> None:
        """Generate Excel workbook with styled summary."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Executive Summary"

        # Title
        ws["A1"] = title
        ws["A1"].font = Font(size=16, bold=True, color="1E293B")
        ws["A2"] = f"Report Type: {report_type.upper()} | Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}"

        # Summary
        ws["A4"] = "Overview"
        ws["A4"].font = Font(bold=True)
        ws["A5"] = data.get("summary", "")

        # KPIs
        kpis = data.get("kpis", {})
        ws["A7"] = "Metric"
        ws["B7"] = "Value"
        ws["A7"].font = Font(bold=True, color="FFFFFF")
        ws["B7"].font = Font(bold=True, color="FFFFFF")
        ws["A7"].fill = PatternFill("solid", fgColor="2563EB")
        ws["B7"].fill = PatternFill("solid", fgColor="2563EB")

        row = 8
        for k, v in kpis.items():
            ws.cell(row=row, column=1, value=str(k))
            ws.cell(row=row, column=2, value=str(v))
            row += 1

        # Insights
        row += 1
        ws.cell(row=row, column=1, value="Strategic Insights").font = Font(bold=True)
        row += 1
        for item in data.get("insights", []):
            ws.cell(row=row, column=1, value=f"- {item}")
            row += 1

        wb.save(str(path))

    def _generate_pptx(
        self, path: Path, title: str, report_type: str, data: dict[str, Any]
    ) -> None:
        """Generate PowerPoint presentation."""
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()

        # Slide 1: Title
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Enterprise Briefing\n{report_type.replace('_', ' ').title()}"

        # Slide 2: Overview & KPIs
        bullet_slide_layout = prs.slide_layouts[1]
        slide2 = prs.slides.add_slide(bullet_slide_layout)
        slide2.shapes.title.text = "Overview & Key Metrics"
        tf2 = slide2.placeholders[1].text_frame
        tf2.text = data.get("summary", "")

        for k, v in data.get("kpis", {}).items():
            p = tf2.add_paragraph()
            p.text = f"{k}: {v}"
            p.level = 1

        # Slide 3: Insights & Next Steps
        slide3 = prs.slides.add_slide(bullet_slide_layout)
        slide3.shapes.title.text = "Strategic Insights & Next Steps"
        tf3 = slide3.placeholders[1].text_frame
        tf3.text = "Key Takeaways:"
        for item in data.get("insights", []):
            p = tf3.add_paragraph()
            p.text = item
            p.level = 1

        prs.save(str(path))


# Global report generator singleton
report_generator = ReportGenerator()
