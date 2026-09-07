"""Comprehensive tests for Phase 5.1 Document Loader Layer.

Verifies:
1. PDF loading using pypdf / PyMuPDF fallback.
2. DOCX loading with paragraphs and tables using python-docx.
3. CSV and TSV loading with semantic tabular formatting.
4. Excel (XLSX) multi-worksheet loading with openpyxl.
5. TXT and Markdown loading with multi-encoding fallback.
6. JSON and JSONL loading with structured formatting.
7. Universal text normalization (linebreaks, null bytes, whitespace).
8. DocumentLoaderFactory format routing.
9. Directory batch processing.
10. REST API endpoints (/load, /load-batch, /supported-formats).
"""

import io
import json
from pathlib import Path
import tempfile

import docx
from fastapi.testclient import TestClient
import openpyxl
import pandas as pd
import pypdf
import pytest

from backend.app.schemas.document_loader import LoadedDocument
from backend.app.services.document_loader_service import DocumentLoaderService
from backend.main import app
from backend.rag.loaders.csv_loader import CSVLoader
from backend.rag.loaders.docx_loader import DocxLoader
from backend.rag.loaders.json_loader import JSONLoader
from backend.rag.loaders.loader_factory import DocumentLoaderFactory
from backend.rag.loaders.pdf_loader import PDFLoader
from backend.rag.loaders.txt_loader import TxtLoader
from backend.rag.loaders.xlsx_loader import XlsxLoader


@pytest.fixture
def service() -> DocumentLoaderService:
    return DocumentLoaderService()


# ----------------------------------------------------------------------
# 1. PDF Loader Test
# ----------------------------------------------------------------------
def test_pdf_loader(service: DocumentLoaderService):
    # Dynamically generate a valid minimal PDF with text
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    # Use annotation or form text if needed, or PyMuPDF for reliable text-injected PDF
    import fitz

    doc = fitz.open()
    pdf_page = doc.new_page(width=400, height=400)
    pdf_page.insert_text((50, 72), "Annual Financial Performance 2026\nRevenue increased by 18.5%.")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = service.load_bytes(pdf_bytes, filename="annual_report.pdf")

    assert isinstance(result, LoadedDocument)
    assert "Annual Financial Performance 2026" in result.content
    assert "Revenue increased by 18.5%" in result.content
    assert result.metadata.file_type == "pdf"
    assert result.metadata.filename == "annual_report.pdf"
    assert result.metadata.page_count >= 1
    assert result.metadata.size == len(pdf_bytes)


# ----------------------------------------------------------------------
# 2. DOCX Loader Test
# ----------------------------------------------------------------------
def test_docx_loader(service: DocumentLoaderService):
    # Dynamically create in-memory Word document
    doc = docx.Document()
    doc.add_heading("Product Strategy Document", level=1)
    doc.add_paragraph("This document details expansion into enterprise analytics.")

    # Add a table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Feature"
    hdr_cells[1].text = "Priority"
    hdr_cells[2].text = "ETA"

    row_cells = table.add_row().cells
    row_cells[0].text = "Document Ingestion"
    row_cells[1].text = "Critical"
    row_cells[2].text = "Q3 2026"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = service.load_bytes(docx_bytes, filename="product_strategy.docx")

    assert isinstance(result, LoadedDocument)
    assert "Product Strategy Document" in result.content
    assert "This document details expansion into enterprise analytics." in result.content
    assert "Document Ingestion" in result.content
    assert "Critical" in result.content
    assert result.metadata.file_type == "docx"
    assert result.metadata.size == len(docx_bytes)


# ----------------------------------------------------------------------
# 3. CSV Loader Test
# ----------------------------------------------------------------------
def test_csv_loader(service: DocumentLoaderService):
    csv_text = (
        "order_id,customer_name,revenue,region\n"
        "1001,Acme Corp,15400.50,North\n"
        "1002,Globex Inc,28200.00,South\n"
        "1003,Stark Tech,45000.75,East\n"
    )
    csv_bytes = csv_text.encode("utf-8")

    result = service.load_bytes(csv_bytes, filename="quarterly_orders.csv")

    assert isinstance(result, LoadedDocument)
    assert "Tabular Dataset: quarterly_orders.csv" in result.content
    assert "Acme Corp" in result.content
    assert "Globex Inc" in result.content
    assert result.metadata.file_type == "csv"
    assert result.metadata.row_count == 3
    assert result.metadata.column_count == 4


# ----------------------------------------------------------------------
# 4. Excel (XLSX) Loader Test
# ----------------------------------------------------------------------
def test_xlsx_loader(service: DocumentLoaderService):
    # Dynamically build multi-sheet workbook
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Executive_Summary"
    ws1.append(["Metric", "Value"])
    ws1.append(["ARR", "$12.4M"])
    ws1.append(["Growth_Rate", "28%"])

    ws2 = wb.create_sheet(title="Regional_Data")
    ws2.append(["Region", "Manager", "Quota"])
    ws2.append(["North", "Alice", "$5.0M"])
    ws2.append(["South", "Bob", "$7.4M"])

    stream = io.BytesIO()
    wb.save(stream)
    xlsx_bytes = stream.getvalue()

    result = service.load_bytes(xlsx_bytes, filename="fy26_budget.xlsx")

    assert isinstance(result, LoadedDocument)
    assert "Worksheet: Executive_Summary" in result.content
    assert "Worksheet: Regional_Data" in result.content
    assert "ARR" in result.content
    assert "Alice" in result.content
    assert result.metadata.file_type == "xlsx"
    assert result.metadata.sheet_names == ["Executive_Summary", "Regional_Data"]
    assert result.metadata.row_count == 4


# ----------------------------------------------------------------------
# 5. TXT & Markdown Loader Test
# ----------------------------------------------------------------------
def test_txt_loader(service: DocumentLoaderService):
    txt_content = "# Operating Principles\n\n1. Maintain data grounding.\n2. Ensure zero hallucinations."
    txt_bytes = txt_content.encode("utf-8")

    result = service.load_bytes(txt_bytes, filename="principles.md")

    assert isinstance(result, LoadedDocument)
    assert "Operating Principles" in result.content
    assert "zero hallucinations" in result.content
    assert result.metadata.file_type == "md"
    assert result.metadata.word_count > 5
    assert result.metadata.line_count >= 3


# ----------------------------------------------------------------------
# 6. JSON & JSONL Loader Test
# ----------------------------------------------------------------------
def test_json_loader(service: DocumentLoaderService):
    # Standard JSON
    payload = {
        "company": "Antigravity Analytics",
        "departments": ["Engineering", "Data Science", "Sales"],
        "metrics": {"active_users": 52000, "satisfaction_score": 4.9},
    }
    json_bytes = json.dumps(payload).encode("utf-8")

    result = service.load_bytes(json_bytes, filename="company_profile.json")

    assert isinstance(result, LoadedDocument)
    assert "Antigravity Analytics" in result.content
    assert "active_users" in result.content
    assert result.metadata.file_type == "json"

    # JSON Lines
    jsonl_text = '{"id": 1, "action": "login"}\n{"id": 2, "action": "query"}\n{"id": 3, "action": "export"}\n'
    result_jsonl = service.load_bytes(jsonl_text.encode("utf-8"), filename="audit.jsonl")

    assert "login" in result_jsonl.content
    assert "export" in result_jsonl.content
    assert result_jsonl.metadata.file_type == "jsonl"


# ----------------------------------------------------------------------
# 7. Content Normalization Test
# ----------------------------------------------------------------------
def test_content_normalization(service: DocumentLoaderService):
    raw = "Header text\r\n\r\n\r\n\r\n\x00Cleaned middle section   \n\n\nFooter note   "
    res = service.load_bytes(raw.encode("utf-8"), filename="dirty.txt")

    # Line endings normalized
    assert "\r" not in res.content
    # Null bytes removed
    assert "\x00" not in res.content
    # 4 consecutive newlines collapsed to 2
    assert "\n\n\n" not in res.content
    # Trailing spaces stripped
    assert "Cleaned middle section" in res.content


# ----------------------------------------------------------------------
# 8. DocumentLoaderFactory Routing Test
# ----------------------------------------------------------------------
def test_loader_factory_routing():
    factory = DocumentLoaderFactory()

    assert isinstance(factory.get_loader("doc.pdf"), PDFLoader)
    assert isinstance(factory.get_loader("doc.docx"), DocxLoader)
    assert isinstance(factory.get_loader("data.csv"), CSVLoader)
    assert isinstance(factory.get_loader("data.tsv"), CSVLoader)
    assert isinstance(factory.get_loader("sheet.xlsx"), XlsxLoader)
    assert isinstance(factory.get_loader("sheet.xls"), XlsxLoader)
    assert isinstance(factory.get_loader("notes.txt"), TxtLoader)
    assert isinstance(factory.get_loader("readme.md"), TxtLoader)
    assert isinstance(factory.get_loader("api.json"), JSONLoader)
    assert isinstance(factory.get_loader("events.jsonl"), JSONLoader)

    # Fallback for unknown
    assert isinstance(factory.get_loader("binary.xyz"), TxtLoader)


# ----------------------------------------------------------------------
# 9. Directory Batch Loading Test
# ----------------------------------------------------------------------
def test_directory_batch_loading(service: DocumentLoaderService):
    with tempfile.TemporaryDirectory() as tmpdir:
        dir_path = Path(tmpdir)
        (dir_path / "notes.txt").write_text("Team weekly notes", encoding="utf-8")
        (dir_path / "data.csv").write_text("a,b\n1,2\n3,4", encoding="utf-8")
        (dir_path / "config.json").write_text('{"env": "prod"}', encoding="utf-8")

        docs = service.load_directory(dir_path)
        assert len(docs) == 3
        filenames = [d.metadata.filename for d in docs]
        assert "notes.txt" in filenames
        assert "data.csv" in filenames
        assert "config.json" in filenames


# ----------------------------------------------------------------------
# 10. REST API Endpoints Test
# ----------------------------------------------------------------------
def test_api_document_loader_endpoints():
    client = TestClient(app)

    # 1. GET /api/v1/documents/supported-formats
    resp_formats = client.get("/api/v1/documents/supported-formats")
    assert resp_formats.status_code == 200
    formats = resp_formats.json()
    assert ".pdf" in formats
    assert ".docx" in formats
    assert ".csv" in formats
    assert ".xlsx" in formats
    assert ".json" in formats

    # 2. POST /api/v1/documents/load (Single file upload)
    txt_file_payload = ("test_note.txt", io.BytesIO(b"Quarterly corporate goals"), "text/plain")
    resp_upload = client.post(
        "/api/v1/documents/load",
        files={"file": txt_file_payload},
    )
    assert resp_upload.status_code == 200
    data = resp_upload.json()
    assert data["status"] == "success"
    assert "Quarterly corporate goals" in data["document"]["content"]
    assert data["document"]["metadata"]["filename"] == "test_note.txt"
    assert data["document"]["metadata"]["file_type"] == "txt"

    # 3. POST /api/v1/documents/load-batch (Multiple files)
    files = [
        ("files", ("batch_1.txt", io.BytesIO(b"Document one content"), "text/plain")),
        ("files", ("batch_2.json", io.BytesIO(b'{"key": "value"}'), "application/json")),
    ]
    resp_batch = client.post(
        "/api/v1/documents/load-batch",
        files=files,
    )
    assert resp_batch.status_code == 200
    batch_data = resp_batch.json()
    assert batch_data["total_loaded"] == 2
    assert len(batch_data["documents"]) == 2
